"""A tiny slide DSL with two renderers, so one description gives an editable .pptx and a pixel-checked PDF.

  * python-pptx writes the .pptx (real text boxes, shapes, a real table, speaker notes).
  * An HTML renderer writes the same slides at 1280x720 (1 inch = 96 px, 1 pt = 4/3 px). Edge renders it
    through Playwright for the PNG previews, the overflow check, and the PDF.

All coordinates are inches on a 13.333 x 7.5 slide. Text boxes have zero inner margin in both renderers.
"""
from __future__ import annotations

import html as _html
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Optional

from PIL import ImageFont

W_IN, H_IN = 13.333, 7.5
LINE_K = 1.22          # Calibri line height / font size, the value PowerPoint and Edge both use
FONT_DIR = Path("C:/Windows/Fonts")


# ----------------------------------------------------------------------------- text model
@dataclass
class Run:
    text: str
    size: float = 14
    bold: bool = False
    italic: bool = False
    color: str = "14232E"
    font: str = "Calibri"


@dataclass
class Para:
    runs: List[Run]
    align: str = "l"
    after: float = 0
    before: float = 0
    lh: float = 1.0
    bullet: bool = False


class R:
    """Override marker used inside P(): R('text', bold=True)."""

    def __init__(self, text, **kw):
        self.text, self.kw = text, kw


def P(*parts, size=14, color="14232E", bold=False, italic=False, font="Calibri", align="l", after=0, before=0, lh=1.0, bullet=False) -> Para:
    runs = []
    for part in parts:
        if isinstance(part, R):
            kw = dict(size=size, color=color, bold=bold, italic=italic, font=font)
            kw.update(part.kw)
            runs.append(Run(part.text, **kw))
        else:
            runs.append(Run(str(part), size=size, color=color, bold=bold, italic=italic, font=font))
    return Para(runs, align=align, after=after, before=before, lh=lh, bullet=bullet)


# ----------------------------------------------------------------------------- elements
@dataclass
class Text:
    x: float; y: float; w: float; h: float
    paras: List[Para]
    valign: str = "t"
    name: str = ""


@dataclass
class Rect:
    x: float; y: float; w: float; h: float
    fill: Optional[str] = None
    line: Optional[str] = None
    lw: float = 1.0
    radius: float = 0.0
    shape: str = "rect"            # rect | oval
    paras: Optional[List[Para]] = None
    valign: str = "m"
    pad: float = 0.12
    name: str = ""


@dataclass
class Line:
    x1: float; y1: float; x2: float; y2: float
    color: str = "52606B"
    lw: float = 1.5
    arrow: bool = False
    dash: bool = False
    name: str = ""


@dataclass
class Img:
    x: float; y: float; w: float; h: float
    path: str
    alt: str = ""


@dataclass
class Table:
    x: float; y: float
    colw: List[float]
    rowh: float
    rows: List[list]               # cell = str | dict(text, bold, color, fill, align, size)
    size: float = 12
    header_fill: str = "0F2B3A"
    header_color: str = "FFFFFF"
    zebra: Optional[str] = "F3F6FA"
    color: str = "14232E"
    name: str = ""


@dataclass
class Slide:
    bg: str = "FFFFFF"
    els: list = field(default_factory=list)
    notes: str = ""
    title: str = ""

    def add(self, *els):
        self.els.extend(els)
        return self


# ----------------------------------------------------------------------------- measurement (for the QA report)
_FONTS = {}


def _font(name, bold, italic, cjk=False):
    key = (name, bold, italic, cjk)
    if key in _FONTS:
        return _FONTS[key]
    if cjk:
        f = ImageFont.truetype(str(FONT_DIR / ("msyhbd.ttc" if bold else "msyh.ttc")), 100)
    elif name == "Cambria":
        f = ImageFont.truetype(str(FONT_DIR / ({(0, 0): "cambria.ttc", (1, 0): "cambriab.ttf", (0, 1): "cambriai.ttf", (1, 1): "cambriaz.ttf"}[(int(bold), int(italic))])), 100)
    else:
        f = ImageFont.truetype(str(FONT_DIR / ({(0, 0): "calibri.ttf", (1, 0): "calibrib.ttf", (0, 1): "calibrii.ttf", (1, 1): "calibriz.ttf"}[(int(bold), int(italic))])), 100)
    _FONTS[key] = f
    return f


def text_width_pt(text, run: Run) -> float:
    w = 0.0
    for ch in text:
        w += _font(run.font, run.bold, run.italic, ord(ch) >= 0x2E80).getlength(ch) * run.size / 100
    return w


def measure_height_in(paras: List[Para], width_in: float) -> float:
    """Wrapped height of a text block in inches (greedy word wrap, like PowerPoint and Edge)."""
    total_pt = 0.0
    for p in paras:
        indent = 0.22 * 72 if p.bullet else 0
        avail = width_in * 72 - indent
        tokens = []                                  # (text, run)
        for r in p.runs:
            for seg_i, seg in enumerate(r.text.split("\n")):
                if seg_i:
                    tokens.append(("\n", r))
                word = ""
                for ch in seg:
                    if ch == " " or ord(ch) >= 0x2E80:
                        if word:
                            tokens.append((word, r))
                            word = ""
                        tokens.append((ch, r))
                    else:
                        word += ch
                if word:
                    tokens.append((word, r))
        lines, cur, size_max, any_text = 1, 0.0, 0.0, False
        line_sizes = []
        for tok, r in tokens:
            if tok == "\n":
                line_sizes.append(size_max or r.size)
                lines += 1; cur = 0.0; size_max = 0.0
                continue
            w = text_width_pt(tok, r)
            if tok == " " and cur == 0.0:
                continue
            if cur + w > avail + 0.01 and tok != " " and cur > 0:
                line_sizes.append(size_max or r.size)
                lines += 1; cur = 0.0; size_max = 0.0
            cur += w
            size_max = max(size_max, r.size)
            any_text = True
        line_sizes.append(size_max or (p.runs[0].size if p.runs else 14))
        total_pt += sum(s * LINE_K * p.lh for s in line_sizes) + p.after + p.before
    return total_pt / 72


# ----------------------------------------------------------------------------- python-pptx renderer
def render_pptx(slides: List[Slide], path: str, title: str, author: str, subject: str = ""):
    from lxml import etree
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_LINE
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    prs.core_properties.title, prs.core_properties.author, prs.core_properties.subject = title, author, subject
    blank = prs.slide_layouts[6]
    AL = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
    AN = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}

    def rgb(h):
        return RGBColor.from_string(h.upper())

    def fill_frame(tf, paras, valign, pad_x=0.0, pad_y=0.0):
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = Inches(pad_x)
        tf.margin_top = tf.margin_bottom = Inches(pad_y)
        tf.vertical_anchor = AN[valign]
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = AL[para.align]
            p.space_after, p.space_before = Pt(para.after), Pt(para.before)
            p.line_spacing = para.lh
            if para.bullet:
                pPr = p._p.get_or_add_pPr()
                pPr.set("marL", str(int(0.22 * 914400)))
                pPr.set("indent", str(-int(0.22 * 914400)))
                bu = etree.SubElement(pPr, qn("a:buFont")); bu.set("typeface", "Arial")
                bc = etree.SubElement(pPr, qn("a:buChar")); bc.set("char", "\u2022")
            for run in para.runs:
                for k, piece in enumerate(run.text.split("\n")):
                    if k:
                        p.add_line_break()
                    r = p.add_run()
                    r.text = piece
                    f = r.font
                    f.size, f.bold, f.italic, f.name = Pt(run.size), run.bold, run.italic, run.font
                    f.color.rgb = rgb(run.color)

    for s in slides:
        sl = prs.slides.add_slide(blank)
        sl.background.fill.solid()
        sl.background.fill.fore_color.rgb = rgb(s.bg)
        for e in s.els:
            if isinstance(e, Text):
                tb = sl.shapes.add_textbox(Inches(e.x), Inches(e.y), Inches(e.w), Inches(e.h))
                tb.name = e.name or "Text"
                fill_frame(tb.text_frame, e.paras, e.valign)
            elif isinstance(e, Rect):
                kind = MSO_SHAPE.OVAL if e.shape == "oval" else (MSO_SHAPE.ROUNDED_RECTANGLE if e.radius > 0 else MSO_SHAPE.RECTANGLE)
                sh = sl.shapes.add_shape(kind, Inches(e.x), Inches(e.y), Inches(e.w), Inches(e.h))
                sh.name = e.name or "Shape"
                sh.shadow.inherit = False
                if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
                    sh.adjustments[0] = min(0.5, e.radius / min(e.w, e.h))
                if e.fill:
                    sh.fill.solid(); sh.fill.fore_color.rgb = rgb(e.fill)
                else:
                    sh.fill.background()
                if e.line:
                    sh.line.color.rgb = rgb(e.line); sh.line.width = Pt(e.lw)
                else:
                    sh.line.fill.background()
                if e.paras:
                    fill_frame(sh.text_frame, e.paras, e.valign, pad_x=e.pad)
            elif isinstance(e, Line):
                c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(e.x1), Inches(e.y1), Inches(e.x2), Inches(e.y2))
                c.name = e.name or "Line"
                c.line.color.rgb = rgb(e.color); c.line.width = Pt(e.lw)
                if e.dash:
                    c.line.dash_style = MSO_LINE.DASH
                if e.arrow:
                    ln = c.line._get_or_add_ln()
                    t = etree.SubElement(ln, qn("a:tailEnd")); t.set("type", "triangle"); t.set("w", "med"); t.set("len", "med")
            elif isinstance(e, Img):
                pic = sl.shapes.add_picture(e.path, Inches(e.x), Inches(e.y), Inches(e.w), Inches(e.h))
                pic._element.nvPicPr.cNvPr.set("descr", e.alt)
                pic.name = "Picture"
            elif isinstance(e, Table):
                nrows, ncols = len(e.rows), len(e.colw)
                gf = sl.shapes.add_table(nrows, ncols, Inches(e.x), Inches(e.y), Inches(sum(e.colw)), Inches(e.rowh * nrows))
                gf.name = e.name or "Table"
                tbl = gf.table
                tbl.first_row = tbl.horz_banding = False
                sid = tbl._tbl.tblPr.find(qn("a:tableStyleId"))
                if sid is None:
                    sid = etree.SubElement(tbl._tbl.tblPr, qn("a:tableStyleId"))
                sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"       # No Style, No Grid
                for ci, cw in enumerate(e.colw):
                    tbl.columns[ci].width = Inches(cw)
                for ri, row in enumerate(e.rows):
                    tbl.rows[ri].height = Inches(e.rowh)
                    for ci, cell in enumerate(row):
                        spec = cell if isinstance(cell, dict) else {"text": str(cell)}
                        header = ri == 0
                        fill = spec.get("fill") or (e.header_fill if header else (e.zebra if (e.zebra and ri % 2 == 0) else "FFFFFF"))
                        c = tbl.cell(ri, ci)
                        c.fill.solid(); c.fill.fore_color.rgb = rgb(fill)
                        c.margin_left = c.margin_right = Inches(0.08)
                        c.margin_top = c.margin_bottom = Inches(0.02)
                        c.vertical_anchor = MSO_ANCHOR.MIDDLE
                        para = P(spec["text"], size=spec.get("size", e.size), bold=spec.get("bold", header),
                                 color=spec.get("color", e.header_color if header else e.color), align=spec.get("align", "l" if ci == 0 else "r"))
                        tf = c.text_frame
                        tf.word_wrap = True
                        p0 = tf.paragraphs[0]
                        p0.alignment = AL[para.align]
                        r = p0.add_run(); r.text = para.runs[0].text
                        r.font.size, r.font.bold, r.font.name = Pt(para.runs[0].size), para.runs[0].bold, "Calibri"
                        r.font.color.rgb = rgb(para.runs[0].color)
        if s.notes:
            sl.notes_slide.notes_text_frame.text = s.notes
    prs.save(path)


# ----------------------------------------------------------------------------- HTML renderer
PX = 96.0


def _px(v):
    return f"{v * PX:.2f}px"


def _run_html(r: Run) -> str:
    txt = _html.escape(r.text).replace("\n", "<br>")
    css = (f"font-size:{r.size}pt;color:#{r.color};font-weight:{'700' if r.bold else '400'};"
           f"font-style:{'italic' if r.italic else 'normal'};font-family:'{r.font}',Calibri,'Microsoft YaHei',sans-serif;")
    return f'<span style="{css}">{txt}</span>'


def _paras_html(paras: List[Para]) -> str:
    out = []
    for p in paras:
        css = f"margin:0;text-align:{ {'l': 'left', 'c': 'center', 'r': 'right'}[p.align] };line-height:{LINE_K * p.lh:.3f};"
        css += f"margin-bottom:{p.after}pt;margin-top:{p.before}pt;"
        cls = ""
        if p.bullet:
            cls = ' class="bul"'
        out.append(f'<p{cls} style="{css}">{"".join(_run_html(r) for r in p.runs)}</p>')
    return "".join(out)


def render_html(slides: List[Slide], path: str, title: str) -> None:
    css = f"""
    @page {{ size: 13.333in 7.5in; margin: 0 }}
    * {{ box-sizing: border-box }}
    html, body {{ margin:0; padding:0; background:#888; }}
    .slide {{ position:relative; width:1280px; height:720px; overflow:hidden; margin:0 auto 24px; page-break-after:always; break-after:page; }}
    @media print {{ html, body {{ background:none }} .slide {{ margin:0 }} }}
    .t, .rt {{ position:absolute; display:flex; flex-direction:column; overflow:visible; }}
    .rt {{ overflow:visible }}
    .abs {{ position:absolute }}
    p.bul {{ padding-left:{0.22*PX:.1f}px; text-indent:-{0.22*PX:.1f}px; }}
    p.bul::before {{ content:"\\2022"; display:inline-block; width:{0.22*PX:.1f}px; text-indent:0; font-family:Arial; }}
    table.tb {{ position:absolute; border-collapse:collapse; table-layout:fixed; }}
    table.tb td {{ padding:0.02in 0.08in; vertical-align:middle; white-space:nowrap; overflow:hidden; }}
    """
    body = []
    for si, s in enumerate(slides):
        parts = [f'<section class="slide" id="s{si + 1}" style="background:#{s.bg}">']
        lines_svg = []
        colors = set()
        for e in s.els:
            if isinstance(e, Text):
                just = {"t": "flex-start", "m": "center", "b": "flex-end"}[e.valign]
                parts.append(f'<div class="t" data-name="{_html.escape(e.name)}" style="left:{_px(e.x)};top:{_px(e.y)};width:{_px(e.w)};height:{_px(e.h)};justify-content:{just}">{_paras_html(e.paras)}</div>')
            elif isinstance(e, Rect):
                st = f"left:{_px(e.x)};top:{_px(e.y)};width:{_px(e.w)};height:{_px(e.h)};"
                st += f"background:#{e.fill};" if e.fill else ""
                st += f"border:{e.lw * 4 / 3:.2f}px solid #{e.line};" if e.line else ""
                st += "border-radius:50%;" if e.shape == "oval" else (f"border-radius:{_px(e.radius)};" if e.radius else "")
                if e.paras:
                    just = {"t": "flex-start", "m": "center", "b": "flex-end"}[e.valign]
                    parts.append(f'<div class="rt" data-name="{_html.escape(e.name)}" style="{st}justify-content:{just};padding:0 {_px(e.pad)}">{_paras_html(e.paras)}</div>')
                else:
                    parts.append(f'<div class="abs" style="{st}"></div>')
            elif isinstance(e, Line):
                colors.add(e.color)
                dash = ' stroke-dasharray="6 5"' if e.dash else ""
                mk = f' marker-end="url(#ah{e.color})"' if e.arrow else ""
                lines_svg.append(f'<line x1="{e.x1 * PX:.2f}" y1="{e.y1 * PX:.2f}" x2="{e.x2 * PX:.2f}" y2="{e.y2 * PX:.2f}" stroke="#{e.color}" stroke-width="{e.lw * 4 / 3:.2f}"{dash}{mk}/>')
            elif isinstance(e, Img):
                parts.append(f'<img class="abs" alt="{_html.escape(e.alt)}" src="{_html.escape(e.path.replace(chr(92), "/"))}" style="left:{_px(e.x)};top:{_px(e.y)};width:{_px(e.w)};height:{_px(e.h)}">')
            elif isinstance(e, Table):
                rows = []
                for ri, row in enumerate(e.rows):
                    tds = []
                    for ci, cell in enumerate(row):
                        spec = cell if isinstance(cell, dict) else {"text": str(cell)}
                        header = ri == 0
                        fill = spec.get("fill") or (e.header_fill if header else (e.zebra if (e.zebra and ri % 2 == 0) else "FFFFFF"))
                        color = spec.get("color", e.header_color if header else e.color)
                        al = {"l": "left", "c": "center", "r": "right"}[spec.get("align", "l" if ci == 0 else "r")]
                        w = "700" if spec.get("bold", header) else "400"
                        tds.append(f'<td style="background:#{fill};color:#{color};font-size:{spec.get("size", e.size)}pt;font-weight:{w};text-align:{al};width:{_px(e.colw[ci])};height:{_px(e.rowh)};font-family:Calibri">{_html.escape(spec["text"])}</td>')
                    rows.append(f"<tr>{''.join(tds)}</tr>")
                parts.append(f'<table class="tb" style="left:{_px(e.x)};top:{_px(e.y)};width:{_px(sum(e.colw))}">{"".join(rows)}</table>')
        if lines_svg:
            defs = "".join(f'<marker id="ah{c}" markerWidth="7" markerHeight="7" refX="6" refY="3.5" orient="auto" markerUnits="userSpaceOnUse"><path d="M0,0 L7,3.5 L0,7 Z" fill="#{c}"/></marker>' for c in colors)
            parts.append(f'<svg class="abs" style="left:0;top:0" width="1280" height="720"><defs>{defs}</defs>{"".join(lines_svg)}</svg>')
        parts.append("</section>")
        body.append("".join(parts))
    Path(path).write_text(f'<!doctype html><html lang="en"><head><meta charset="utf-8"><title>{_html.escape(title)}</title><style>{css}</style></head><body>{"".join(body)}</body></html>', encoding="utf-8")


# ----------------------------------------------------------------------------- export + QA through Edge
def export_html(html_path: str, png_dir: str, pdf_path: str):
    """Render with the system Edge: per-slide PNG, the PDF, and a list of text boxes whose content overflows."""
    from playwright.sync_api import sync_playwright

    png_dir_p = Path(png_dir)
    png_dir_p.mkdir(parents=True, exist_ok=True)
    problems = []
    with sync_playwright() as p:
        b = p.chromium.launch(channel="msedge", headless=True)
        ctx = b.new_context(viewport={"width": 1280, "height": 720}, device_scale_factor=2)
        page = ctx.new_page()
        page.goto(Path(html_path).resolve().as_uri())
        page.wait_for_load_state("networkidle")
        page.evaluate("document.fonts.ready")
        n = page.locator("section.slide").count()
        for i in range(n):
            sec = page.locator(f"#s{i + 1}")
            sec.screenshot(path=str(png_dir_p / f"slide_{i + 1:02d}.png"))
        res = page.evaluate("""() => {
            const out = [];
            document.querySelectorAll('section.slide').forEach((sec, i) => {
              const sb = sec.getBoundingClientRect();
              sec.querySelectorAll('.t, .rt').forEach(el => {
                const r = el.getBoundingClientRect();
                // content height: the sum of the paragraph boxes, wherever they overflow to
                let top = Infinity, bottom = -Infinity, right = -Infinity;
                el.querySelectorAll('p').forEach(pp => { const q = pp.getBoundingClientRect(); top = Math.min(top, q.top); bottom = Math.max(bottom, q.bottom); right = Math.max(right, q.right); });
                const over = (bottom - r.bottom > 1.5) || (r.top - top > 1.5);
                const wide = el.scrollWidth - el.clientWidth > 2;
                const out_of_slide = r.left < sb.left - 1 || r.top < sb.top - 1 || r.right > sb.right + 1 || r.bottom > sb.bottom + 1;
                if (over || wide || out_of_slide) out.push({slide: i + 1, name: el.dataset.name || '', text: el.innerText.slice(0, 50), over, wide, out_of_slide, boxH: Math.round(r.height), contentH: Math.round(bottom - top)});
              });
            });
            return out;
        }""")
        problems = res
        pdf_page = ctx.new_page()
        pdf_page.goto(Path(html_path).resolve().as_uri())
        pdf_page.wait_for_load_state("networkidle")
        pdf_page.pdf(path=pdf_path, width="13.333in", height="7.5in", print_background=True, margin={"top": "0", "right": "0", "bottom": "0", "left": "0"})
        b.close()
    return problems
